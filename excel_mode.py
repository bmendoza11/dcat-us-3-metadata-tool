import io
import json
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter


# Excel Mode intentionally excludes Additional DCAT-US Properties.
COLUMNS = [
    (
        'bureau',
        'Which bureau are you submitting metadata for?',
        'Dropdown',
        'Select a bureau.',
        True
    ),
    (
        'catalog_contact',
        'Who should be contacted with questions about this catalog?',
        'Dropdown',
        'Select from contacts.json for the bureau.',
        True
    ),
    (
        'dataset_title',
        'What is the dataset called?',
        'Text',
        'Example: Current Population Survey',
        True
    ),
    (
        'dataset_description',
        'Dataset description in plain language.',
        'Long text',
        'Describe what the dataset contains.',
        True
    ),
    (
        'identifier',
        'Dataset identifier',
        'Auto-assigned',
        'Automatically assigned by the app (e.g., BEA-000001).',
        True
    ),
    (
        'publishing_office',
        'Which office publishes these data?',
        'Dropdown',
        'Select an office or enter a new office name.',
        True
    ),
    (
        'dataset_contact',
        'What is the best contact for questions about this dataset?',
        'Dropdown',
        'Select from contacts.json for the bureau.',
        True
    ),
    (
        'keywords',
        'What words would someone use to search for these data?',
        'Semicolon-separated text',
        'Example: employment; labor; wages',
        True
    ),
    (
        'themes',
        'What theme does this dataset belong to?',
        'Semicolon-separated text',
        'Example: labor; economic conditions',
        True
    ),
    (
        'theme_descriptions',
        'Descriptions for any new themes',
        'Semicolon-separated key=value',
        'Example: new theme=Description of the theme',
        False
    ),
    (
        'access_rights',
        'Who is allowed to access these data?',
        'Dropdown',
        'These data are public / restricted / not public.',
        True
    ),
    (
        'access_restriction',
        'Are there any restrictions on getting these data?',
        'Dropdown',
        'No or Yes.',
        True
    ),
    (
        'access_status',
        'Access restriction status',
        'Dropdown',
        'Only used when access_restriction=Yes.',
        False
    ),
    (
        'specific_access_restrictions',
        'Specific access restrictions',
        'Semicolon-separated text',
        'Example: FOIA (b)(3) Statute; Other',
        False
    ),
    (
        'access_restriction_note',
        'Additional information about the access restriction',
        'Long text',
        'Optional note.',
        False
    ),
    (
        'cui',
        'Does the dataset contain Controlled Unclassified Information (CUI)?',
        'Dropdown',
        'No or Yes.',
        True
    ),
    (
        'cui_banner',
        'CUI Banner Marking',
        'Dropdown',
        'CONTROLLED, SP-PROPIN, SP-CTI, SP-PRVCY, SP-PII, SP-HLTH, SP-FIN.',
        False
    ),
    (
        'cui_designation',
        'Which agency designated the information as CUI?',
        'Long text',
        'Include contact information when possible.',
        False
    ),
    (
        'use_restriction',
        'Are there other rules about how these data may be used?',
        'Dropdown',
        'No or Yes.',
        True
    ),
    (
        'use_restriction_types',
        'Types of use rule',
        'Semicolon-separated text',
        'Use restriction; License; Rights',
        False
    ),
    (
        'use_status',
        'Use restriction status',
        'Dropdown',
        'Only used when Use restriction is selected.',
        False
    ),
    (
        'specific_use_restrictions',
        'Specific use restrictions',
        'Semicolon-separated text',
        'Copyright; Donor Restrictions; Public Law; Other',
        False
    ),
    (
        'use_restriction_note',
        'Use restriction note',
        'Long text',
        'Optional note.',
        False
    ),
    (
        'license',
        'What license applies to these data?',
        'Text',
        'Required if License is selected.',
        False
    ),
    (
        'rights',
        'What other rights have not been covered?',
        'Long text',
        'Required if Rights is selected.',
        False
    ),
    (
        'temporal_start',
        'Start date',
        'Date text',
        'YYYY, YYYY-MM, or YYYY-MM-DD',
        True
    ),
    (
        'temporal_end',
        'End date',
        'Date text',
        'YYYY, YYYY-MM, or YYYY-MM-DD',
        True
    ),
    (
        'spatial_granularity',
        'Spatial Granularity',
        'Dropdown',
        'Continent, Country, State, County, City, ZIP Code, Census Tract, Other.',
        False
    ),
    (
        'spatial_other',
        'Other spatial granularity',
        'Text',
        'Only used when spatial_granularity=Other.',
        False
    ),
    (
        'geographic_coverage',
        'Where do these data cover?',
        'Text',
        'Example: United States, Washington, DC, or worldwide',
        True
    ),
    (
        'modified',
        'When were these data last changed?',
        'Date text',
        'YYYY, YYYY-MM, or YYYY-MM-DD',
        True
    ),
    (
        'has_dictionary',
        'Does this dataset have a data dictionary?',
        'Dropdown',
        'No or Yes.',
        False
    ),
    (
        'dictionary_url',
        'Data dictionary URL',
        'URL',
        'Required if has_dictionary=Yes.',
        False
    ),
    (
        'dictionary_format',
        'Data dictionary format',
        'Text',
        'Example: HTML, PDF, CSV, XLSX',
        False
    ),
    (
        'has_landing_page',
        'Is there a webpage/landing page to get the data from?',
        'Dropdown',
        'No or Yes.',
        False
    ),
    (
        'landing_page',
        'Landing page URL',
        'URL',
        'Required if has_landing_page=Yes.',
        False
    ),
    (
        'contract_number',
        'Contract number by which these data were acquired',
        'Text',
        'Optional; maximum 40 characters.',
        False
    ),
]


CHOICES = {
    "access_rights": [
        "These data are public",
        "These data are restricted",
        "These data are not public",
    ],
    "yes_no": ["No", "Yes"],
    "access_status": [
        "Restricted - Fully",
        "Restricted - Partly",
        "Restricted - Possibly",
        "Undetermined",
        "Unrestricted",
    ],
    "cui_banner": [
        "CONTROLLED",
        "SP-PROPIN",
        "SP-CTI",
        "SP-PRVCY",
        "SP-PII",
        "SP-HLTH",
        "SP-FIN",
    ],
    "spatial_granularity": [
        "Continent",
        "Country",
        "State",
        "County",
        "City",
        "ZIP Code",
        "Census Tract",
        "Other",
    ],
}


def split_values(value):
    if value is None:
        return []

    if isinstance(value, list):
        return [
            str(item).strip()
            for item in value
            if str(item).strip()
        ]

    return [
        item.strip()
        for item in re.split(r"[;\n]", str(value))
        if item.strip()
    ]


def parse_theme_descriptions(value):
    result = {}

    for item in split_values(value):
        if "=" not in item:
            continue

        name, description = item.split("=", 1)

        if name.strip():
            result[name.strip().lower()] = description.strip()

    return result


def safe_range_name(text):
    name = re.sub(
        r"[^A-Za-z0-9_]",
        "_",
        text
    ).strip("_") or "contacts"

    if name[0].isdigit():
        name = "_" + name

    return name[:240]


def contact_display(contact):
    name = contact.get("name", "Unnamed contact")
    organization = contact.get("organization", "")
    email = contact.get("email", "")

    if organization and email:
        return f"{name} — {organization} — {email}"

    if organization:
        return f"{name} — {organization}"

    return name


def generalized_contact(contacts):
    if not contacts:
        return None

    preferred_terms = [
        "call center",
        "contact center",
        "general information",
        "information center",
        "information services",
        "webmaster",
        "general contact",
        "customer service",
        "customer contact",
        "general",
        "contact",
        "public affairs",
    ]

    ranked = []

    for index, contact in enumerate(contacts):
        searchable = " ".join(
            str(contact.get(field, ""))
            for field in (
                "name",
                "email",
                "organization",
            )
        ).lower()

        score = next(
            (
                position
                for position, term in enumerate(preferred_terms)
                if term in searchable
            ),
            999,
        )

        ranked.append(
            (score, index, contact)
        )

    return sorted(
        ranked,
        key=lambda item: (item[0], item[1])
    )[0][2]


def _contact_for_name(name, contacts):
    target = str(name or "").strip().lower()

    for contact in contacts:
        if (
            contact_display(contact).strip().lower() == target
            or contact.get("name", "").strip().lower() == target
        ):
            return contact

    return None


def catalog_to_rows(catalog, bureaus, contacts_by_bureau):
    if not catalog:
        return []

    publisher = (
        (catalog.get("publisher") or {}).get("name")
        or ""
    )

    bureau = next(
        (
            name
            for name in bureaus
            if name.lower() == publisher.lower()
        ),
        None,
    )

    if not bureau:
        bureau = next(
            (
                name
                for name in bureaus
                if name.lower() in publisher.lower()
                or publisher.lower() in name.lower()
            ),
            bureaus[0],
        )

    catalog_contact = ""
    contact_points = catalog.get("contactPoint") or []

    if contact_points:
        catalog_name = contact_points[0].get("fn", "")
        match = _contact_for_name(
            catalog_name,
            contacts_by_bureau.get(bureau, []),
        )
        catalog_contact = (
            contact_display(match)
            if match
            else catalog_name
        )

    rows = []

    for index, dataset in enumerate(
        catalog.get("dataset", []) or [],
        start=1,
    ):
        contact_points = dataset.get("contactPoint") or [{}]
        dataset_contact = contact_points[0].get("fn", "")
        contact_match = _contact_for_name(
            dataset_contact,
            contacts_by_bureau.get(bureau, []),
        )

        if contact_match:
            dataset_contact = contact_display(contact_match)

        keywords = dataset.get("keyword", []) or []
        themes = []

        for theme in dataset.get("theme", []) or []:
            if isinstance(theme, dict):
                themes.append(theme.get("prefLabel", ""))
            else:
                themes.append(str(theme))

        publisher_obj = dataset.get("publisher") or {}
        temporal = (dataset.get("temporal") or [{}])[0]
        spatial = (dataset.get("spatial") or [{}])[0]

        rows.append(
            {
                "bureau": bureau,
                "catalog_contact": catalog_contact,
                "dataset_title": dataset.get("title", ""),
                "dataset_description": dataset.get("description", ""),
                "identifier": dataset.get(
                    "identifier",
                    f"{bureaus[0]}-{index:06d}",
                ),
                "publishing_office": publisher_obj.get(
                    "name",
                    "",
                ),
                "dataset_contact": dataset_contact,
                "keywords": "; ".join(keywords),
                "themes": "; ".join(themes),
                "access_rights": dataset.get(
                    "accessRights",
                    "",
                ),
                "access_restriction": (
                    "Yes"
                    if dataset.get("accessRestriction")
                    else "No"
                ),
                "cui": (
                    "Yes"
                    if dataset.get("cuiRestriction")
                    else "No"
                ),
                "use_restriction": (
                    "Yes"
                    if (
                        dataset.get("useRestriction")
                        or dataset.get("license")
                        or dataset.get("rights")
                    )
                    else "No"
                ),
                "temporal_start": temporal.get(
                    "startDate",
                    "",
                ),
                "temporal_end": temporal.get(
                    "endDate",
                    "",
                ),
                "spatial_granularity": spatial.get(
                    "spatialGranularity",
                    "",
                ),
                "geographic_coverage": spatial.get(
                    "prefLabel",
                    "",
                ),
                "modified": dataset.get(
                    "modified",
                    "",
                ),
                "contract_number": dataset.get(
                    "contractNumber",
                    "",
                ),
            }
        )

    return rows


def create_template(
    bureaus,
    contacts_by_bureau,
    offices_by_bureau,
    tags,
    themes,
    existing_catalog=None,
    default_bureau=None,
):
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Data Entry"

    dictionary = workbook.create_sheet("Data Dictionary")
    instructions = workbook.create_sheet("Instructions")
    references = workbook.create_sheet("Reference Lists")
    new_contacts = workbook.create_sheet("New Contacts")

    header_fill = PatternFill(
        "solid",
        fgColor="005EA8",
    )
    header_font = Font(
        color="FFFFFF",
        bold=True,
    )
    thin = Side(
        style="thin",
        color="D9E2F3",
    )

    for column_index, (key, *_rest) in enumerate(
        COLUMNS,
        start=1,
    ):
        cell = worksheet.cell(
            1,
            column_index,
            key,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            wrap_text=True,
        )
        cell.border = Border(
            bottom=thin,
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(
            max(len(key) + 4, 16),
            34,
        )

    rows = catalog_to_rows(
        existing_catalog,
        bureaus,
        contacts_by_bureau,
    ) if existing_catalog else []

    if not rows:
        rows = [{}]

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, *_rest) in enumerate(
            COLUMNS,
            start=1,
        ):
            worksheet.cell(
                row_index,
                column_index,
                row.get(key, ""),
            )

    if default_bureau and not existing_catalog:
        bureau_column = next(
            index
            for index, column in enumerate(COLUMNS, start=1)
            if column[0] == "bureau"
        )
        catalog_contact_column = next(
            index
            for index, column in enumerate(COLUMNS, start=1)
            if column[0] == "catalog_contact"
        )
        dataset_contact_column = next(
            index
            for index, column in enumerate(COLUMNS, start=1)
            if column[0] == "dataset_contact"
        )

        worksheet.cell(
            2,
            bureau_column,
            default_bureau,
        )

        default_contact = generalized_contact(
            contacts_by_bureau.get(
                default_bureau,
                [],
            )
        )

        if default_contact:
            display = contact_display(default_contact)
            worksheet.cell(
                2,
                catalog_contact_column,
                display,
            )
            worksheet.cell(
                2,
                dataset_contact_column,
                display,
            )

    dictionary.append(
        [
            "Column name",
            "Full question",
            "Expected/example response",
            "Input type",
            "Required",
        ]
    )

    for cell in dictionary[1]:
        cell.fill = header_fill
        cell.font = header_font

    for key, question, input_type, example, required in COLUMNS:
        dictionary.append(
            [
                key,
                question,
                example,
                input_type,
                "Required" if required else "Optional",
            ]
        )

    dictionary.freeze_panes = "A2"

    for column, width in {
        "A": 26,
        "B": 60,
        "C": 50,
        "D": 24,
        "E": 14,
    }.items():
        dictionary.column_dimensions[column].width = width

    for row in dictionary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    instructions.append(["DCAT-US 3.0 Excel Mode"])
    instructions.append([
        "1. Fill out one row per dataset in Data Entry."
    ])
    instructions.append([
        "2. Keep the short column names unchanged; use Data Dictionary for the full questions."
    ])
    instructions.append([
        "3. Contacts are dropdowns sourced from contacts.json for the selected bureau."
    ])
    instructions.append([
        "4. To add contacts to the database, use the New Contacts sheet."
    ])
    instructions.append([
        "5. Keywords and themes are separated with semicolons. New keywords are added to tags.json."
    ])
    instructions.append([
        "6. New themes use name=description in theme_descriptions and are saved to themes.json."
    ])
    instructions.append([
        "7. Additional DCAT-US Properties are excluded from Excel Mode; Guided Mode still contains them."
    ])

    instructions.column_dimensions["A"].width = 120

    for row in instructions.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    references.append(
        ["Bureaus"] + list(bureaus)
    )
    references.append(
        ["Access Rights"] + CHOICES["access_rights"]
    )
    references.append(
        ["Yes / No"] + CHOICES["yes_no"]
    )
    references.append(
        ["Access Status"] + CHOICES["access_status"]
    )
    references.append(
        ["CUI Banner"] + CHOICES["cui_banner"]
    )
    references.append(
        ["Spatial Granularity"]
        + CHOICES["spatial_granularity"]
    )
    references.append(
        [
            "Use Restriction Types",
            "Use restriction",
            "License",
            "Rights",
        ]
    )
    references.append(
        [
            "Specific Use",
            "Copyright",
            "Donor Restrictions",
            "Public Law",
            "Other",
        ]
    )
    references.append(
        [
            "Specific Access",
            "FOIA (b)(1) National Security",
            "FOIA (b)(2) Internal Personnel Rules and Practices",
            "FOIA (b)(3) Statute",
            "FOIA (b)(4) Trade Secrets and Commercial Information",
            "FOIA (b)(5) Personal Information",
            "FOIA (b)(7) Law Enforcement",
            "FOIA (b)(8) Financial Institutions",
            "FOIA (b)(9) Geological and Geophysical Information",
            "Other",
        ]
    )
    references.append(
        ["Tags"] + list(tags)
    )
    references.append(
        ["Themes"] + [
            theme.get("name", "")
            for theme in themes
        ]
    )

    for index, bureau in enumerate(bureaus, start=1):
        range_name = (
            safe_range_name(bureau)
            + "_contacts"
        )
        values = [
            contact_display(contact)
            for contact in contacts_by_bureau.get(
                bureau,
                [],
            )
        ]
        start_column = 12 + ((index - 1) * 2)

        references.cell(
            1,
            start_column,
            bureau,
        )

        for row_index, value in enumerate(
            values,
            start=2,
        ):
            references.cell(
                row_index,
                start_column,
                value,
            )

        if values:
            reference = (
                f"'Reference Lists'!"
                f"${get_column_letter(start_column)}$2:"
                f"${get_column_letter(start_column)}${len(values) + 1}"
            )
            workbook.defined_names.add(
                DefinedName(
                    range_name,
                    attr_text=reference,
                )
            )

    references.sheet_state = "hidden"

    for column_index, heading in enumerate(
        [
            "bureau",
            "name",
            "email",
            "phone",
            "organization",
        ],
        start=1,
    ):
        cell = new_contacts.cell(
            1,
            column_index,
            heading,
        )
        cell.fill = header_fill
        cell.font = header_font
        new_contacts.column_dimensions[
            get_column_letter(column_index)
        ].width = 30

    new_contacts.append([
        default_bureau or "",
        "",
        "",
        "",
        "",
    ])

    def add_list_validation(
        column_key,
        formula,
        max_row=500,
    ):
        column = next(
            index
            for index, item in enumerate(COLUMNS, start=1)
            if item[0] == column_key
        )

        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )

        worksheet.add_data_validation(validation)
        validation.add(
            f"{get_column_letter(column)}2:"
            f"{get_column_letter(column)}{max_row}"
        )

    add_list_validation(
        "bureau",
        "'Reference Lists'!$B$1:$L$1",
    )
    add_list_validation(
        "access_rights",
        "'Reference Lists'!$B$2:$D$2",
    )
    add_list_validation(
        "access_restriction",
        "'Reference Lists'!$B$3:$C$3",
    )
    add_list_validation(
        "access_status",
        "'Reference Lists'!$B$4:$F$4",
    )
    add_list_validation(
        "cui",
        "'Reference Lists'!$B$3:$C$3",
    )
    add_list_validation(
        "cui_banner",
        "'Reference Lists'!$B$5:$H$5",
    )
    add_list_validation(
        "spatial_granularity",
        "'Reference Lists'!$B$6:$I$6",
    )
    add_list_validation(
        "has_dictionary",
        "'Reference Lists'!$B$3:$C$3",
    )
    add_list_validation(
        "has_landing_page",
        "'Reference Lists'!$B$3:$C$3",
    )

    for column_key in (
        "catalog_contact",
        "dataset_contact",
    ):
        column = next(
            index
            for index, item in enumerate(COLUMNS, start=1)
            if item[0] == column_key
        )

        validation = DataValidation(
            type="list",
            formula1=(
                "=INDIRECT(IFERROR("
                "VLOOKUP($A2,"
                "'Reference Lists'!$A$1:$A$11,"
                "2,FALSE),\"\"))"
            ),
            allow_blank=True,
        )

        worksheet.add_data_validation(validation)
        validation.add(
            f"{get_column_letter(column)}2:"
            f"{get_column_letter(column)}500"
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(COLUMNS))}"
        f"{max(2, len(rows) + 1)}"
    )

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()


def read_rows(workbook):
    worksheet = workbook["Data Entry"]
    headers = [cell.value for cell in worksheet[1]]
    rows = []

    for values in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not any(
            value not in (None, "")
            for value in values
        ):
            continue

        rows.append(
            {
                header: values[index]
                if index < len(values)
                else ""
                for index, header in enumerate(headers)
            }
        )

    return rows


def read_new_contacts(workbook):
    if "New Contacts" not in workbook.sheetnames:
        return []

    worksheet = workbook["New Contacts"]
    contacts = []

    for values in worksheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        values = list(values) + [""] * 5
        bureau, name, email, phone, organization = values[:5]

        if bureau and name:
            contacts.append(
                {
                    "bureau": str(bureau).strip(),
                    "name": str(name).strip(),
                    "email": str(email or "").strip(),
                    "phone": str(phone or "").strip(),
                    "organization": str(organization or "").strip(),
                }
            )

    return contacts
