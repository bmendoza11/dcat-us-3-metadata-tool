import json
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


# ============================================================
# EXCEL FIELD DEFINITIONS
# ============================================================
#
# These are the fields represented in the Excel workbook.
#
# "field"     = stable machine-readable Excel column name
# "question"  = human-readable question
# "type"      = how the Excel value should be entered
# "required"  = whether the value is required
# "options"   = dropdown options, if applicable
# "dcat"      = DCAT-US property
#
# Keep these names stable once users start receiving workbooks.
# ============================================================

EXCEL_FIELDS = [
    {
        "field": "dataset_title",
        "question": "What is the dataset called?",
        "type": "text",
        "required": True,
        "options": [],
        "dcat": "title",
    },
    {
        "field": "dataset_description",
        "question": "Dataset description in plain language.",
        "type": "long_text",
        "required": True,
        "options": [],
        "dcat": "description",
    },
    {
        "field": "publisher_office",
        "question": "Which office publishes these data?",
        "type": "dropdown",
        "required": True,
        "options": [],
        "list_name": "OFFICES",
        "dcat": "publisher",
    },
    {
        "field": "contact_name",
        "question": (
            "What is the best contact for questions "
            "about this dataset?"
        ),
        "type": "text",
        "required": True,
        "options": [],
        "dcat": "contactPoint.fn",
    },
    {
        "field": "contact_email",
        "question": "Dataset contact email.",
        "type": "email",
        "required": False,
        "options": [],
        "dcat": "contactPoint.hasEmail",
    },
    {
        "field": "contact_phone",
        "question": "Dataset contact phone.",
        "type": "text",
        "required": False,
        "options": [],
        "dcat": "contactPoint.hasTelephone",
    },
    {
        "field": "contact_organization",
        "question": "Dataset contact organization.",
        "type": "text",
        "required": False,
        "options": [],
        "dcat": "contactPoint.organization",
    },
    {
        "field": "keywords",
        "question": (
            "What words would someone use to search "
            "for these data?"
        ),
        "type": "list",
        "required": True,
        "options": [],
        "dcat": "keyword",
        "help": (
            "Separate multiple keywords with semicolons. "
            "Example: employment; labor; unemployment"
        ),
    },
    {
        "field": "themes",
        "question": (
            "What theme does this dataset belong to?"
        ),
        "type": "list",
        "required": True,
        "options": [],
        "list_name": "THEMES",
        "dcat": "theme",
        "help": (
            "Separate multiple themes with semicolons."
        ),
    },
    {
        "field": "access_rights",
        "question": (
            "Who is allowed to access these data?"
        ),
        "type": "dropdown",
        "required": True,
        "options": [
            "These data are public",
            "These data are restricted",
            "These data are not public",
        ],
        "dcat": "accessRights",
    },
    {
        "field": "has_access_restriction",
        "question": (
            "Are there any restrictions on getting these data?"
        ),
        "type": "dropdown",
        "required": True,
        "options": [
            "No",
            "Yes",
        ],
        "dcat": "accessRestriction",
    },
    {
        "field": "access_status",
        "question": "What is the access restriction status?",
        "type": "dropdown",
        "required": False,
        "options": [
            "Restricted - Fully",
            "Restricted - Partly",
            "Restricted - Possibly",
            "Undetermined",
            "Unrestricted",
        ],
        "dcat": "accessRestriction.restrictionStatus",
    },
    {
        "field": "specific_access",
        "question": (
            "What is the specific access restriction?"
        ),
        "type": "list",
        "required": False,
        "options": [
            "FOIA (b)(1) National Security",
            "FOIA (b)(2) Internal Personnel Rules and Practices",
            "FOIA (b)(3) Statute",
            "FOIA (b)(4) Trade Secrets and Commercial Information",
            "FOIA (b)(5) Privileged Inter-Agency or Intra-Agency Information",
            "FOIA (b)(6) Personal Information",
            "FOIA (b)(7) Law Enforcement",
            "FOIA (b)(8) Financial Institutions",
            "FOIA (b)(9) Geological and Geophysical Information",
            "Other",
        ],
        "dcat": "accessRestriction.specificRestriction",
    },
    {
        "field": "access_note",
        "question": (
            "Additional information about the access restriction."
        ),
        "type": "long_text",
        "required": False,
        "options": [],
        "dcat": "accessRestriction.restrictionNote",
    },
    {
        "field": "has_cui",
        "question": (
            "Does the dataset contain Controlled "
            "Unclassified Information (CUI)?"
        ),
        "type": "dropdown",
        "required": True,
        "options": [
            "No",
            "Yes",
        ],
        "dcat": "cuiRestriction",
    },
    {
        "field": "cui_banner",
        "question": (
            "Which CUI Banner Marking are these data "
            "associated with?"
        ),
        "type": "dropdown",
        "required": False,
        "options": [
            "CONTROLLED",
            "SP-PROPIN",
            "SP-CTI",
            "SP-PRVCY",
            "SP-PII",
            "SP-HLTH",
            "SP-FIN",
        ],
        "dcat": "cuiRestriction.cuiBannerMarking",
    },
    {
        "field": "cui_designation",
        "question": (
            "Which agency designated the information as CUI?"
        ),
        "type": "long_text",
        "required": False,
        "options": [],
        "dcat": "cuiRestriction.designationIndicator",
    },
    {
        "field": "has_use_restriction",
        "question": (
            "Are there any other rules about how these data "
            "may be used?"
        ),
        "type": "dropdown",
        "required": True,
        "options": [
            "No",
            "Yes",
        ],
        "dcat": "useRestriction",
    },
    {
        "field": "restriction_types",
        "question": (
            "What type of rule applies?"
        ),
        "type": "list",
        "required": False,
        "options": [
            "Use restriction",
            "License",
            "Rights",
        ],
        "dcat": "useRestriction",
    },
    {
        "field": "use_status",
        "question": "Use restriction status.",
        "type": "dropdown",
        "required": False,
        "options": [
            "Restricted - Fully",
            "Restricted - Partly",
            "Restricted - Possibly",
            "Undetermined",
            "Unrestricted",
        ],
        "dcat": "useRestriction.restrictionStatus",
    },
    {
        "field": "specific_use",
        "question": "Specific use restriction.",
        "type": "list",
        "required": False,
        "options": [
            "Copyright",
            "Donor Restrictions",
            "Public Law",
            "Other",
        ],
        "dcat": "useRestriction.specificRestriction",
    },
    {
        "field": "use_note",
        "question": "Use restriction note.",
        "type": "long_text",
        "required": False,
        "options": [],
        "dcat": "useRestriction.restrictionNote",
    },
    {
        "field": "license",
        "question": (
            "What license applies to these data?"
        ),
        "type": "text",
        "required": False,
        "options": [],
        "dcat": "license",
    },
    {
        "field": "rights",
        "question": (
            "What other rights have not been covered?"
        ),
        "type": "long_text",
        "required": False,
        "options": [],
        "dcat": "rights",
    },
    {
        "field": "temporal_start",
        "question": (
            "What time period do these data cover? "
            "Start date."
        ),
        "type": "date",
        "required": True,
        "options": [],
        "dcat": "temporal.start",
        "help": (
            "Use YYYY, YYYY-MM, or YYYY-MM-DD."
        ),
    },
    {
        "field": "temporal_end",
        "question": (
            "What time period do these data cover? "
            "End date."
        ),
        "type": "date",
        "required": True,
        "options": [],
        "dcat": "temporal.end",
        "help": (
            "Use YYYY, YYYY-MM, or YYYY-MM-DD."
        ),
    },
    {
        "field": "spatial_granularity",
        "question": "Spatial granularity.",
        "type": "dropdown",
        "required": False,
        "options": [
            "",
            "Continent",
            "Country",
            "State",
            "County",
            "City",
            "ZIP Code",
            "Census Tract",
            "Other",
        ],
        "dcat": "spatial.granularity",
    },
    {
        "field": "spatial",
        "question": (
            "Where do these data cover?"
        ),
        "type": "text",
        "required": True,
        "options": [],
        "dcat": "spatial",
    },
    {
        "field": "modified",
        "question": (
            "When were these data last changed?"
        ),
        "type": "date",
        "required": True,
        "options": [],
        "dcat": "modified",
        "help": (
            "Use YYYY, YYYY-MM, or YYYY-MM-DD."
        ),
    },
    {
        "field": "has_dictionary",
        "question": (
            "Does this dataset have a data dictionary?"
        ),
        "type": "dropdown",
        "required": False,
        "options": [
            "No",
            "Yes",
        ],
        "dcat": "describedBy",
    },
    {
        "field": "dictionary_url",
        "question": "Data dictionary URL.",
        "type": "url",
        "required": False,
        "options": [],
        "dcat": "describedBy.url",
    },
    {
        "field": "dictionary_format",
        "question": "Data dictionary format.",
        "type": "text",
        "required": False,
        "options": [],
        "dcat": "describedBy.format",
    },
    {
        "field": "has_landing_page",
        "question": (
            "Is there a webpage/landing page to get "
            "the data from?"
        ),
        "type": "dropdown",
        "required": False,
        "options": [
            "No",
            "Yes",
        ],
        "dcat": "landingPage",
    },
    {
        "field": "landing_page",
        "question": "Landing page URL.",
        "type": "url",
        "required": False,
        "options": [],
        "dcat": "landingPage",
    },
    {
        "field": "contract_number",
        "question": (
            "Contract number by which these data were acquired."
        ),
        "type": "text",
        "required": False,
        "options": [],
        "dcat": "contractNumber",
    },
]


# ============================================================
# UTILITY FUNCTIONS
# ============================================================

def clean_excel_value(value):
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

    return str(value).strip()


def split_excel_list(value):
    """
    Convert:

        employment; labor; unemployment

    into:

        ["employment", "labor", "unemployment"]
    """

    value = clean_excel_value(value)

    if not value:
        return []

    return [
        item.strip()
        for item in value.split(";")
        if item.strip()
    ]


def join_excel_list(value):
    """
    Convert a JSON list into an Excel-friendly string.
    """

    if not value:
        return ""

    if isinstance(value, list):
        return "; ".join(
            str(item)
            for item in value
            if item is not None
        )

    return str(value)


def parse_json_cell(value, field_name):
    value = clean_excel_value(value)

    if not value:
        return None

    try:
        return json.loads(value)

    except json.JSONDecodeError as error:
        raise ValueError(
            f"{field_name} contains invalid JSON: "
            f"{error.msg}"
        )


def is_valid_dcat_date(value):
    value = clean_excel_value(value)

    if not value:
        return True

    if re.fullmatch(r"\d{4}", value):
        return True

    if re.fullmatch(r"\d{4}-\d{2}", value):
        year, month = map(
            int,
            value.split("-")
        )

        return 1 <= month <= 12

    if re.fullmatch(
        r"\d{4}-\d{2}-\d{2}",
        value
    ):
        from datetime import date

        try:
            date.fromisoformat(value)
            return True
        except ValueError:
            return False

    return False


# ============================================================
# WORKBOOK CREATION
# ============================================================

def create_excel_template(
    bureau_info,
    offices=None,
    themes=None,
    existing_rows=None,
    existing_catalog_contact=None,
    additional_property_definitions=None,
):
    """
    Create a complete Excel workbook.

    Sheets:

    Instructions
    Catalog
    Datasets
    Codebook
    Lists
    """

    wb = Workbook()

    instructions_ws = wb.active
    instructions_ws.title = "Instructions"

    catalog_ws = wb.create_sheet("Catalog")
    datasets_ws = wb.create_sheet("Datasets")
    codebook_ws = wb.create_sheet("Codebook")
    lists_ws = wb.create_sheet("Lists")

    additional_property_definitions = (
        additional_property_definitions or {}
    )

    # --------------------------------------------------------
    # INSTRUCTIONS
    # --------------------------------------------------------

    instructions_ws["A1"] = (
        "DCAT-US 3.0 Metadata Builder"
    )

    instructions_ws["A1"].font = Font(
        size=18,
        bold=True,
        color="005EA8"
    )

    instructions_ws["A3"] = "Bureau"
    instructions_ws["B3"] = bureau_info.get(
        "publisher",
        ""
    )

    instructions_ws["A5"] = (
        "How to use this workbook"
    )

    instructions_ws["A5"].font = Font(
        bold=True,
        color="005EA8"
    )

    instructions = [
        "Complete the Catalog sheet with catalog-level information.",
        "Complete one row on the Datasets sheet for each dataset.",
        "Do not rename columns on the Datasets sheet.",
        "Use the Codebook sheet to understand each field.",
        "Dropdown fields contain controlled choices.",
        "For list fields, separate multiple values with semicolons.",
        "For JSON fields, enter valid JSON.",
        "Dates must use YYYY, YYYY-MM, or YYYY-MM-DD.",
        "Do not modify the Codebook or Lists sheets.",
        "Save the workbook as XLSX when finished.",
        "Upload the completed workbook back into the application.",
    ]

    for row_number, instruction in enumerate(
        instructions,
        start=6
    ):
        instructions_ws.cell(
            row=row_number,
            column=1,
            value=f"• {instruction}"
        )

    instructions_ws.column_dimensions[
        "A"
    ].width = 100

    # --------------------------------------------------------
    # CATALOG SHEET
    # --------------------------------------------------------

    catalog_ws["A1"] = "Catalog Information"

    catalog_ws["A1"].font = Font(
        size=16,
        bold=True,
        color="005EA8"
    )

    catalog_fields = [
        (
            "bureau",
            "Bureau",
            bureau_info.get("publisher", "")
        ),
        (
            "catalog_contact_name",
            "Catalog contact name",
            ""
        ),
        (
            "catalog_contact_email",
            "Catalog contact email",
            ""
        ),
        (
            "catalog_contact_phone",
            "Catalog contact phone",
            ""
        ),
        (
            "catalog_contact_organization",
            "Catalog contact organization",
            bureau_info.get("publisher", "")
        ),
    ]

    if existing_catalog_contact:
        catalog_fields = [
            (
                field,
                label,
                existing_catalog_contact.get(
                    field,
                    default_value
                )
            )
            for field, label, default_value
            in catalog_fields
        ]

    for row_number, (
        field,
        label,
        value
    ) in enumerate(
        catalog_fields,
        start=3
    ):

        catalog_ws.cell(
            row=row_number,
            column=1,
            value=field
        )

        catalog_ws.cell(
            row=row_number,
            column=2,
            value=label
        )

        catalog_ws.cell(
            row=row_number,
            column=3,
            value=value
        )

        catalog_ws.cell(
            row=row_number,
            column=1
        ).font = Font(bold=True)

    catalog_ws["A10"] = (
        "Only edit values in column C."
    )

    catalog_ws["A10"].font = Font(
        italic=True,
        color="666666"
    )

    # --------------------------------------------------------
    # DATASET HEADERS
    # --------------------------------------------------------

    dataset_headers = (
        ["dataset_number"]
        + [
            field["field"]
            for field in EXCEL_FIELDS
        ]
        + list(additional_property_definitions.keys())
    )

    header_fill = PatternFill(
        "solid",
        fgColor="005EA8"
    )

    for column_number, header in enumerate(
        dataset_headers,
        start=1
    ):

        cell = datasets_ws.cell(
            row=1,
            column=column_number,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

    datasets_ws.freeze_panes = "A2"

    datasets_ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(dataset_headers))}1"
    )

    # --------------------------------------------------------
    # CODEBOOK
    # --------------------------------------------------------

    codebook_headers = [
        "field",
        "question",
        "type",
        "required",
        "DCAT-US property",
        "options",
        "help",
    ]

    for column_number, header in enumerate(
        codebook_headers,
        start=1
    ):

        cell = codebook_ws.cell(
            row=1,
            column=column_number,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

    codebook_row = 2

    for field in EXCEL_FIELDS:

        options = field.get(
            "options",
            []
        )

        if options:
            options_text = " | ".join(
                str(option)
                for option in options
            )
        elif field.get("list_name"):
            options_text = (
                f"See Lists sheet: "
